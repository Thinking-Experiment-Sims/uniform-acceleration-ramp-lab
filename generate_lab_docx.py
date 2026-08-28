import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# Colors matching The Thinking Experiment palette
HEX_TEAL = "0F7E9B"
HEX_AMBER = "D67B19"
HEX_LIGHT_BG = "E6F4F8"
HEX_LIGHT_AMBER = "FDF3E8"
HEX_BORDER = "CBD5E1"
HEX_DARK = "1E293B"

RGB_TEAL = RGBColor(15, 126, 155)
RGB_AMBER = RGBColor(214, 123, 25)
RGB_DARK = RGBColor(30, 41, 59)
RGB_MUTED = RGBColor(100, 116, 139)

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color=HEX_BORDER, sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_doc_header(doc, title, subtitle):
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(2)
    p_sub.paragraph_format.space_before = Pt(0)
    run_sub = p_sub.add_run("THE THINKING EXPERIMENT | MODELING INSTRUCTION")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(8.5)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGB_AMBER

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(3)
    p_title.paragraph_format.space_before = Pt(0)
    run_title = p_title.add_run(title)
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(16)
    run_title.font.bold = True
    run_title.font.color.rgb = RGB_TEAL

    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_after = Pt(10)
    p_desc.paragraph_format.space_before = Pt(0)
    run_desc = p_desc.add_run(subtitle)
    run_desc.font.name = 'Arial'
    run_desc.font.size = Pt(9.5)
    run_desc.font.color.rgb = RGB_MUTED

def add_scenario_section(doc, scenario_num, title, setup_desc, page_break_before=False):
    if page_break_before:
        doc.add_page_break()

    # Section Header Table
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(7.0)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, HEX_TEAL)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run_t = p.add_run(f"{scenario_num}. {title}")
    run_t.font.name = 'Arial'
    run_t.font.size = Pt(11.5)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(255, 255, 255)

    # Question a: Observation
    p_a = doc.add_paragraph()
    p_a.paragraph_format.space_before = Pt(6)
    p_a.paragraph_format.space_after = Pt(4)
    run_a = p_a.add_run(f"a. {setup_desc}")
    run_a.font.name = 'Arial'
    run_a.font.size = Pt(9.5)

    # Question b: Motion Map Box
    tbl_mm = doc.add_table(rows=2, cols=1)
    tbl_mm.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_mm.autofit = False
    tbl_mm.columns[0].width = Inches(7.0)
    set_table_borders(tbl_mm, HEX_BORDER, sz="6")

    c_mm_hdr = tbl_mm.cell(0, 0)
    set_cell_background(c_mm_hdr, HEX_LIGHT_BG)
    set_cell_margins(c_mm_hdr, top=60, bottom=60, left=100, right=100)
    p_hdr = c_mm_hdr.paragraphs[0]
    p_hdr.paragraph_format.space_after = Pt(0)
    run_h = p_hdr.add_run("b. Draw a motion map. Include velocity and acceleration vectors.")
    run_h.font.name = 'Arial'
    run_h.font.size = Pt(9)
    run_h.font.bold = True
    run_h.font.color.rgb = RGB_TEAL

    c_mm_body = tbl_mm.cell(1, 0)
    set_cell_background(c_mm_body, "FFFFFF")
    set_cell_margins(c_mm_body, top=140, bottom=140, left=100, right=100)
    p_body = c_mm_body.paragraphs[0]
    p_body.paragraph_format.space_after = Pt(0)
    p_body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_track = p_body.add_run("0 position ───────────────────────────────────────────────────────────── +")
    run_track.font.name = 'Courier New'
    run_track.font.size = Pt(9)
    run_track.font.color.rgb = RGB_MUTED

    # Questions c & d Table
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    tbl_cd = doc.add_table(rows=1, cols=2)
    tbl_cd.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_cd.autofit = False
    tbl_cd.columns[0].width = Inches(3.45)
    tbl_cd.columns[1].width = Inches(3.45)
    set_table_borders(tbl_cd, HEX_BORDER, sz="4")

    # c
    c1 = tbl_cd.cell(0, 0)
    set_cell_margins(c1, top=60, bottom=60, left=80, right=80)
    p_c = c1.paragraphs[0]
    p_c.paragraph_format.space_after = Pt(0)
    if scenario_num == 5:
        r_c = p_c.add_run("c. Does the direction of velocity change?\n[   ] Yes   [   ] No")
    else:
        r_c = p_c.add_run("c. Is velocity positive or negative?\n[   ] Positive   [   ] Negative")
    r_c.font.name = 'Arial'
    r_c.font.size = Pt(9)
    r_c.font.bold = True

    # d
    c2 = tbl_cd.cell(0, 1)
    set_cell_margins(c2, top=60, bottom=60, left=80, right=80)
    p_d = c2.paragraphs[0]
    p_d.paragraph_format.space_after = Pt(0)
    if scenario_num == 5:
        r_d = p_d.add_run("d. Does the direction of acceleration change?\n[   ] Yes   [   ] No")
    else:
        r_d = p_d.add_run("d. Is acceleration positive or negative?\n[   ] Positive   [   ] Negative")
    r_d.font.name = 'Arial'
    r_d.font.size = Pt(9)
    r_d.font.bold = True

    # Graphs & Analysis Section (e, f, g, h)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    tbl_efgh = doc.add_table(rows=4, cols=3)
    tbl_efgh.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_efgh.autofit = False
    tbl_efgh.columns[0].width = Inches(2.2)
    tbl_efgh.columns[1].width = Inches(2.2)
    tbl_efgh.columns[2].width = Inches(2.6)
    set_table_borders(tbl_efgh, HEX_BORDER, sz="4")

    # Header Row
    headers = [
        "e. Predict Graphs",
        "f. Recorded Motion Detector",
        "Slope Analysis"
    ]
    for idx, text in enumerate(headers):
        cell_h = tbl_efgh.cell(0, idx)
        set_cell_background(cell_h, HEX_LIGHT_BG)
        set_cell_margins(cell_h, top=60, bottom=60, left=60, right=60)
        p_h = cell_h.paragraphs[0]
        p_h.paragraph_format.space_after = Pt(0)
        run_th = p_h.add_run(text)
        run_th.font.name = 'Arial'
        run_th.font.size = Pt(8.5)
        run_th.font.bold = True
        run_th.font.color.rgb = RGB_TEAL

    # Row 1: x vs t
    r1_c0 = tbl_efgh.cell(1, 0)
    set_cell_margins(r1_c0, top=40, bottom=40, left=40, right=40)
    r1_c0.paragraphs[0].add_run("Position (x vs t)\n\n+ | \n  | \n0 └─── t").font.size = Pt(7.5)

    r1_c1 = tbl_efgh.cell(1, 1)
    set_cell_margins(r1_c1, top=40, bottom=40, left=40, right=40)
    r1_c1.paragraphs[0].add_run("Position (x vs t)\n\n+ | \n  | \n0 └─── t").font.size = Pt(7.5)

    r1_c2 = tbl_efgh.cell(1, 2)
    set_cell_margins(r1_c2, top=40, bottom=40, left=40, right=40)
    r1_c2.paragraphs[0].add_run(
        "g. The slope of position-time is:\n"
        "( constant / increasing / decreasing )\n"
        "and ( positive / negative )\n"
        "and represents ___________________."
    ).font.size = Pt(8)

    # Row 2: v vs t
    r2_c0 = tbl_efgh.cell(2, 0)
    set_cell_margins(r2_c0, top=40, bottom=40, left=40, right=40)
    r2_c0.paragraphs[0].add_run("Velocity (v vs t)\n\n+ | \n0 ┼─── t\n- |").font.size = Pt(7.5)

    r2_c1 = tbl_efgh.cell(2, 1)
    set_cell_margins(r2_c1, top=40, bottom=40, left=40, right=40)
    r2_c1.paragraphs[0].add_run("Velocity (v vs t)\n\n+ | \n0 ┼─── t\n- |").font.size = Pt(7.5)

    r2_c2 = tbl_efgh.cell(2, 2)
    set_cell_margins(r2_c2, top=40, bottom=40, left=40, right=40)
    r2_c2.paragraphs[0].add_run(
        "h. The slope of velocity-time is:\n"
        "( constant / increasing / decreasing )\n"
        "and ( positive / negative )\n"
        "and represents ___________________."
    ).font.size = Pt(8)

    # Row 3: a vs t
    r3_c0 = tbl_efgh.cell(3, 0)
    set_cell_margins(r3_c0, top=40, bottom=40, left=40, right=40)
    r3_c0.paragraphs[0].add_run("Acceleration (a vs t)\n\n+ | \n0 ┼─── t\n- |").font.size = Pt(7.5)

    r3_c1 = tbl_efgh.cell(3, 1)
    set_cell_margins(r3_c1, top=40, bottom=40, left=40, right=40)
    r3_c1.paragraphs[0].add_run("Acceleration (a vs t)\n\n+ | \n0 ┼─── t\n- |").font.size = Pt(7.5)

    r3_c2 = tbl_efgh.cell(3, 2)
    set_cell_margins(r3_c2, top=40, bottom=40, left=40, right=40)
    r3_c2.paragraphs[0].add_run(
        "Acceleration is constant:\n"
        "a = g·sin(θ)\n"
        "Direction: [  ] Positive  [  ] Negative"
    ).font.size = Pt(8)

def generate_handout():
    doc = Document()
    
    # Page setup (0.5 in margins for maximum usable space)
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    format_doc_header(
        doc,
        "Uniformly Accelerated Particle Model",
        "Lab Extension: Increasing and Decreasing Speed | Name: _______________________ Date: ________ Period: ____"
    )

    # Scenario 1
    add_scenario_section(
        doc, 1, "Increasing speed in the positive direction",
        "Without using the motion detector, observe the motion of the cart as it starts from rest and rolls down the incline away from the detector (+ direction).",
        page_break_before=False
    )

    # Scenario 2
    add_scenario_section(
        doc, 2, "Decreasing speed in the positive direction",
        "Without using the motion detector, observe the motion of the cart slowing after an initial push up the ramp away from the detector. Answer questions for the cart while coasting.",
        page_break_before=True
    )

    # Scenario 3
    add_scenario_section(
        doc, 3, "Increasing speed in the negative direction",
        "Observe the motion of the cart starting from rest at the top (+ position) and rolling down the incline towards the motion detector (0 position).",
        page_break_before=True
    )

    # Scenario 4
    add_scenario_section(
        doc, 4, "Decreasing speed in the negative direction",
        "Observe the motion of the cart slowing after an initial push up the ramp towards the motion detector (0 position). Answer questions for the cart while coasting.",
        page_break_before=True
    )

    # Scenario 5
    add_scenario_section(
        doc, 5, "Up and down the ramp (The Turnaround)",
        "Observe the motion of the cart after an initial push up the ramp. The cart moves away (+), stops momentarily at the apex, turns around, and rolls back down towards the detector.",
        page_break_before=True
    )

    out_path = "Uniform_Acceleration_Student_Handout.docx"
    doc.save(out_path)
    print(f"Generated {out_path}")

if __name__ == "__main__":
    generate_handout()
